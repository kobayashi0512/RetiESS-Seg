"""Build a complete patient-level AFIO manifest for later 2D extension work.

All 49 unique paired records in the released archive are indexed.  The one
duplicate spreadsheet row is represented in the provenance audit, but does not
create a fictitious second image example.
"""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from PIL import Image
from openpyxl import load_workbook

from audit_afio import clean_name, consensus_label


def size(path: Path) -> list[int]:
    with Image.open(path) as image:
        return [image.width, image.height]


def main() -> None:
    root = Path("data/afio_glaucoma/extracted")
    raw = [path for path in root.rglob("*.jpg") if not path.name.endswith(("_ILM.jpg", "_RPE.jpg"))]
    raw_by_stem = {clean_name(path.stem): path for path in raw}
    wb = load_workbook(next(root.rglob("*.xlsx")), data_only=True, read_only=True)
    unique = {}
    duplicates = []
    for values in wb.active.iter_rows(min_row=3, values_only=True):
        color_stem, subject = clean_name(values[0]), str(values[11] or "").strip().replace(" ", "")
        if not color_stem or not subject:
            continue
        if color_stem in unique:
            duplicates.append(color_stem)
            continue
        bscan_stem = color_stem.replace("_Color_", "_B-scan_")
        color, bscan = raw_by_stem[color_stem], raw_by_stem[bscan_stem]
        unique[color_stem] = {
            "patient_id": subject,
            "color_path": str(color.resolve()),
            "oct_path": str(bscan.resolve()),
            "color_size": size(color),
            "oct_size": size(bscan),
            "label": consensus_label([values[2], values[4], values[6], values[8]]),
            "label_votes": [str(value).strip().lower() if value is not None else None for value in (values[2], values[4], values[6], values[8])],
            "ilm_mask_path": str((bscan.with_name(bscan.stem + "_ILM.jpg")).resolve()) if bscan.with_name(bscan.stem + "_ILM.jpg").exists() else None,
            "rpe_mask_path": str((bscan.with_name(bscan.stem + "_RPE.jpg")).resolve()) if bscan.with_name(bscan.stem + "_RPE.jpg").exists() else None,
        }
    records = [unique[key] for key in sorted(unique)]
    if len(records) != 49:
        raise RuntimeError(f"Expected complete 49-image AFIO manifest, got {len(records)}")
    report = {
        "dataset": "AFIO complete public archive",
        "record_count": len(records),
        "patient_count": len({item["patient_id"] for item in records}),
        "labels": dict(Counter(item["label"] for item in records)),
        "records_with_ilm_mask": sum(item["ilm_mask_path"] is not None for item in records),
        "records_with_rpe_mask": sum(item["rpe_mask_path"] is not None for item in records),
        "duplicate_spreadsheet_stems": duplicates,
        "records": records,
    }
    destination = Path("results/afio_manifest.json")
    destination.write_text(json.dumps(report, indent=2), encoding="utf-8")
    with Path("results/afio_manifest.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=[
            "patient_id", "color_path", "oct_path", "color_size", "oct_size", "label",
            "ilm_mask_path", "rpe_mask_path",
        ], extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)
    print(json.dumps({key: value for key, value in report.items() if key != "records"}, indent=2))


if __name__ == "__main__":
    main()
